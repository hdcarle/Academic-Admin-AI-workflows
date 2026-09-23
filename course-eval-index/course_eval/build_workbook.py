"""Build the Excel workbook from data/sections.json + config.yaml (+ optional verbatims.csv).

Every summary number is a live formula over the 'Course Index' sheet, except the top-box
block (computed once from raw response counts and labeled STATIC).
"""
import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .core import DEFAULT_SEASON_ORDER, load_yaml

FONT = "Arial"
NAVY, BLUE, GREEN, GRAY, PURP = "1F3864", "2F5597", "548235", "595959", "6B4C9A"
PALETTE = {"blue": BLUE, "green": GREEN, "gray": GRAY, "purple": PURP, "tan": "7F6000", "navy": NAVY}
thin = Side(style="thin", color="BFBFBF")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(b=False, i=False, c="000000", s=10):
    return Font(name=FONT, bold=b, italic=i, color=c, size=s)


def fill(h):
    return PatternFill("solid", start_color=h, end_color=h)


def note_for(x, cfg, small_n):
    out = []
    for rule in cfg.get("notes", []):
        w = rule.get("when", {})
        ok = True
        if "year" in w and x["year"] != w["year"]:
            ok = False
        if "season" in w and x["season"] != w["season"]:
            ok = False
        if "season_in" in w and x["season"] not in w["season_in"]:
            ok = False
        if "survey" in w and x["survey"] != w["survey"]:
            ok = False
        if "format" in w and x["format"] != w["format"]:
            ok = False
        if ok:
            out.append(rule["text"].format(year=x["year"], season=x["season"]))
    if len(x["section_codes"]) > 1:
        out.append("Combined report (multiple sections on one report); # taught is the combined enrollment.")
    if x["responded"] is not None and x["responded"] < small_n:
        out.append(f"Small sample (fewer than {small_n} respondents); counted in summary averages, weighted by respondents.")
    return " ".join(out)


def build(cfg, sections, out_path, verbatims=None, quiet=False):
    inst = cfg.get("instructor", {})
    surveys = cfg.get("surveys") or [{"id": "default", "label": "All"}]
    survey_label = {s["id"]: s.get("label", s["id"]) for s in surveys}
    groups = cfg["groups"]
    gmap = {g["id"]: g for g in groups}
    items = cfg["items"]
    scfg = cfg.get("summary", {})
    small_n = scfg.get("small_sample_threshold", 5)
    min_resp = scfg.get("min_respondents", 1)
    summ_survey = scfg.get("avg_overall_survey")  # survey id or None (= all sections)
    titles = cfg.get("course_titles", {})
    muted = set(cfg.get("muted_when_online", []))
    recs = []
    for x in sections:
        y = dict(x)
        y["term_label"] = f"{x['season']} {x['year']}"
        recs.append(y)
    n_rows = len(recs)
    if not n_rows:
        raise SystemExit("No sections found. Run `parse` first and check data/warnings.txt.")

    # ---- column plan
    primary = [g["id"] for g in groups if g.get("primary")] or [groups[0]["id"]]
    other = [g["id"] for g in groups if g["id"] not in primary]
    cols_of = {gid: [it for it in items if it["group"] == gid] for gid in gmap}
    col = {}
    c = 6
    for gid in primary:
        for it in cols_of[gid]:
            col[it["key"]] = c
            c += 1
    c_over = c
    c_sec, c_resp, c_rr, c_survey, c_inc, c_year, c_notes, c_title, c_name, c_src = range(c_over + 1, c_over + 11)
    c = c_src + 1
    for gid in other:
        for it in cols_of[gid]:
            col[it["key"]] = c
            c += 1
    last_col = c - 1
    HR, R0 = 4, 5
    R1 = R0 + n_rows - 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Course Index"
    who = inst.get("display_name", "Instructor")
    ws["A1"] = f"Teaching Evaluation Index — {cfg.get('institution', '')} ({who})".replace(" ()", "")
    ws["A1"].font = f(True, c=NAVY, s=14)
    avg_groups = [g["id"] for g in groups if g.get("in_avg_overall")]
    ws["A2"] = ("One row per course section. Ratings are the section means printed on each evaluation report. "
                "Avg. Overall = simple average of the teaching-item means flagged for it in config.yaml "
                "(evaluations do not print a composite). See 'Method & Notes' for definitions.")
    ws["A2"].font = f(i=True, c=GRAY, s=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(8, c_over))
    ws.row_dimensions[2].height = 38

    def band(c1, c2, text, color):
        ws.cell(row=3, column=c1, value=text)
        for cc in range(c1, c2 + 1):
            cell = ws.cell(row=3, column=cc)
            cell.fill = fill(color)
            cell.font = f(True, c="FFFFFF")
            cell.alignment = Alignment(horizontal="left", vertical="center")

    band(1, 5, "COURSE", NAVY)
    for gid in primary + other:
        its = cols_of[gid]
        if its:
            g = gmap[gid]
            band(col[its[0]["key"]], col[its[-1]["key"]], g.get("title", gid), PALETTE.get(g.get("color", "blue"), BLUE))
    band(c_over, c_over, "OVERALL", GREEN)
    band(c_sec, c_src, "SECTION DETAIL & SOURCE", GRAY)

    heads = {1: ("Course", 12), 2: ("Format", 11), 3: ("Term", 13), 4: ("UG or GR", 9), 5: ("# taught", 9)}
    for it in items:
        heads[col[it["key"]]] = (it["label"], 15)
    heads[c_over] = ("Avg. Overall", 11)
    heads[c_sec] = ("Section", 10)
    heads[c_resp] = ("Respondents", 12)
    heads[c_rr] = ("Response rate", 10)
    heads[c_survey] = ("Survey version", 11)
    heads[c_inc] = ("Counted in summary averages? (1=yes)", 13)
    heads[c_year] = ("Year", 7)
    heads[c_notes] = ("Notes", 60)
    heads[c_title] = ("Course title", 30)
    heads[c_name] = ("Name on report", 14)
    heads[c_src] = ("Source file", 40)
    for cc, (h, w) in heads.items():
        cell = ws.cell(row=HR, column=cc, value=h)
        cell.font = f(True, s=9)
        cell.fill = fill("D9E1F2")
        cell.border = BORD
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[L(cc)].width = w
    ws.row_dimensions[HR].height = 118
    ws.row_dimensions[3].height = 20

    def rng_group_row(gid, r):
        its = cols_of[gid]
        return f"{L(col[its[0]['key']])}{r}:{L(col[its[-1]['key']])}{r}" if its else None

    avg_ranges_for = lambda r: [rng_group_row(g, r) for g in avg_groups if cols_of[g]]
    numeric_cols = [col[it["key"]] for it in items] + [c_over]
    left_cols = (1, 2, 3, 4, c_sec, c_survey, c_notes, c_title, c_name, c_src)
    for ri, x in enumerate(recs):
        r = R0 + ri
        for cc, v in enumerate([x["course"], x["format"], x["term_label"], x["level"], x["enrolled"]], 1):
            ws.cell(row=r, column=cc, value=v)
        for k, v in x["items"].items():
            if k in col:
                ws.cell(row=r, column=col[k], value=v["mean"])
        rr_ = avg_ranges_for(r)
        ws.cell(row=r, column=c_over, value=f"=IFERROR(AVERAGE({','.join(rr_)}),0)")
        ws.cell(row=r, column=c_sec, value=x["section"])
        ws.cell(row=r, column=c_resp, value=x["responded"])
        ws.cell(row=r, column=c_rr, value=f'=IFERROR({L(c_resp)}{r}/E{r},"")')
        ws.cell(row=r, column=c_survey, value=survey_label.get(x["survey"], x["survey"]))
        ws.cell(row=r, column=c_inc, value=f"=IF(AND({L(c_resp)}{r}>={min_resp},COUNT({','.join(rr_)})>0),1,0)")
        ws.cell(row=r, column=c_year, value=x["year"])
        note = note_for(x, cfg, small_n)
        ws.cell(row=r, column=c_notes, value=note)
        ws.cell(row=r, column=c_title, value=titles.get(x["course"], ""))
        ws.cell(row=r, column=c_name, value=x.get("name_on_report") or "")
        ws.cell(row=r, column=c_src, value=x["file"])
        for cc in range(1, last_col + 1):
            cell = ws.cell(row=r, column=cc)
            cell.font = f()
            cell.border = BORD
            cell.alignment = Alignment(vertical="top", wrap_text=(cc == c_notes),
                                       horizontal="left" if cc in left_cols else "center")
        for cc in numeric_cols:
            ws.cell(row=r, column=cc).number_format = "0.00"
        oc = ws.cell(row=r, column=c_over)
        oc.font, oc.fill = f(True), fill("E2EFDA")
        ws.cell(row=r, column=c_rr).number_format = "0%"
        if x["format"] == "Online":
            for k in muted:
                if k in col:
                    ws.cell(row=r, column=col[k]).font = f(i=True, c="7F7F7F")
        if x["responded"] is not None and x["responded"] < small_n:
            for cc in (c_resp, c_notes):
                ws.cell(row=r, column=cc).fill = fill("FFF2CC")
        if summ_survey and x["survey"] != summ_survey:
            ws.cell(row=r, column=c_survey).fill = fill("FCE4D6")
        ws.row_dimensions[r].height = 30 if len(note) > 60 else 16

    # ---- totals block under the table
    T0 = R1 + 2
    rngE, rngB, rngD = (f"${c}${R0}:${c}${R1}" for c in ("E", "B", "D"))
    ws.cell(row=T0, column=1, value="STUDENTS TAUGHT (sum of enrollment, all sections in this index)").font = f(True, c=NAVY, s=11)
    ws.cell(row=T0 + 1, column=1, value="Group").font = f(True, s=9)
    ws.cell(row=T0 + 1, column=5, value="# taught").font = f(True, s=9)
    ws.cell(row=T0 + 1, column=6, value="Sections").font = f(True, s=9)
    tot = [("All sections", f"=SUM({rngE})", f"=COUNTA($A${R0}:$A${R1})"),
           ("In-person", f'=SUMIFS({rngE},{rngB},"In-person")', f'=COUNTIFS({rngB},"In-person")'),
           ("Online", f'=SUMIFS({rngE},{rngB},"Online")', f'=COUNTIFS({rngB},"Online")'),
           ("Undergraduate (UG)", f'=SUMIFS({rngE},{rngD},"UG")', f'=COUNTIFS({rngD},"UG")'),
           ("Graduate (GR)", f'=SUMIFS({rngE},{rngD},"GR")', f'=COUNTIFS({rngD},"GR")')]
    for i, (lab, fs, fc) in enumerate(tot):
        r = T0 + 2 + i
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=5, value=fs)
        ws.cell(row=r, column=6, value=fc)
        for cc in range(1, 7):
            cell = ws.cell(row=r, column=cc)
            cell.font, cell.fill, cell.border = f(True), fill("F2F2F2"), BORD
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="center")
    ws.freeze_panes = ws.cell(row=R0, column=6)
    ws.auto_filter.ref = f"A{HR}:{L(last_col)}{R1}"
    for gid in other:
        g = gmap[gid]
        if g.get("collapse") and cols_of[gid]:
            ws.column_dimensions.group(L(col[cols_of[gid][0]["key"]]), L(col[cols_of[gid][-1]["key"]]),
                                       hidden=False, outline_level=1)
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.tabColor = NAVY

    # ================= Summary =================
    S = wb.create_sheet("Summary")
    S.sheet_properties.tabColor = GREEN
    CI = "'Course Index'!"
    rg = lambda cc: f"{CI}${L(cc)}${R0}:${L(cc)}${R1}"
    E_, B_, D_ = rg(5), rg(2), rg(4)
    OV, RS, SV, INC, YR, CRS = rg(c_over), rg(c_resp), rg(c_survey), rg(c_inc), rg(c_year), rg(1)
    ALL = f"ISNUMBER({E_})"
    SF = f'({SV}="{survey_label[summ_survey]}")' if summ_survey else "1"
    S["A1"] = "Teaching Summary"
    S["A1"].font = f(True, c=NAVY, s=14)
    scope = (f"the '{survey_label[summ_survey]}' survey version" if summ_survey else "all sections")
    S["A2"] = ("Every figure on this sheet is a live formula over the 'Course Index' sheet, except the block marked static. "
               f"Rating averages are respondent-weighted and use {scope}; sections and students are counted for every section.")
    S["A2"].font = f(i=True, c=GRAY, s=9)
    S.merge_cells("A2:H2")
    S["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    S.row_dimensions[2].height = 30

    def hdr(r, labels):
        for i, t in enumerate(labels, 1):
            cell = S.cell(row=r, column=i, value=t)
            cell.font, cell.fill, cell.border = f(True, c="FFFFFF"), fill(BLUE), BORD
            cell.alignment = Alignment(horizontal="center" if i > 1 else "left", vertical="center", wrap_text=True)

    def put(r, vals, fmts=None, bold=False):
        for i, v in enumerate(vals, 1):
            cell = S.cell(row=r, column=i, value=v)
            cell.font, cell.border = f(bold), BORD
            cell.alignment = Alignment(horizontal="left" if i == 1 else "center")
            if fmts and fmts[i - 1]:
                cell.number_format = fmts[i - 1]

    def cnt(cond, sf=True):
        return f"=SUMPRODUCT({cond}*{SF if sf else 1}*{INC})"

    def rsp(cond):
        return f"=SUMPRODUCT({cond}*{SF}*{INC}*{RS})"

    def wavg(cond):
        return (f'=IFERROR(SUMPRODUCT({cond}*{SF}*{INC}*{OV}*{RS})/SUMPRODUCT({cond}*{SF}*{INC}*{RS}),"n/a")')

    r = 4
    S.cell(row=r, column=1, value="1. Students taught (enrollment)").font = f(True, c=NAVY, s=12)
    hdr(r + 1, ["Group", "Sections", "# taught", "% of all students"])
    grp = [("All sections", ALL), ("In-person", f'({B_}="In-person")'), ("Online", f'({B_}="Online")'),
           ("Undergraduate (UG)", f'({D_}="UG")'), ("Graduate (GR)", f'({D_}="GR")')]
    for i, (lab, cond) in enumerate(grp):
        rr = r + 2 + i
        put(rr, [lab, f"=SUMPRODUCT({cond}*1)", f"=SUMPRODUCT({cond}*{E_})", f"=C{rr}/$C${r + 2}"],
            [None, "0", "#,##0", "0.0%"], bold=(i == 0))
    r = r + 2 + len(grp) + 1
    S.cell(row=r, column=1, value="Format × level (# taught)").font = f(True)
    hdr(r + 1, ["", "UG", "GR", "Total"])
    for i, fm in enumerate(["In-person", "Online"]):
        rr = r + 2 + i
        put(rr, [fm, f'=SUMIFS({E_},{B_},"{fm}",{D_},"UG")', f'=SUMIFS({E_},{B_},"{fm}",{D_},"GR")', f"=B{rr}+C{rr}"],
            [None, "#,##0", "#,##0", "#,##0"])
    rr = r + 4
    put(rr, ["Total", f"=B{r + 2}+B{r + 3}", f"=C{r + 2}+C{r + 3}", f"=D{r + 2}+D{r + 3}"],
        [None, "#,##0", "#,##0", "#,##0"], bold=True)
    r = rr + 2

    S.cell(row=r, column=1, value=f"2. Average Overall (5-point scale) — {scope}, respondent-weighted").font = f(True, c=NAVY, s=12)
    hdr(r + 1, ["Group", "Sections counted", "Respondents", "Avg. Overall"])
    grp2 = [("All rated sections", ALL)] + grp[1:]
    for i, (lab, cond) in enumerate(grp2):
        rr = r + 2 + i
        put(rr, [lab, cnt(cond), rsp(cond), wavg(cond)], [None, "0", "#,##0", "0.00"], bold=(i == 0))
    r = r + 2 + len(grp2) + 1

    S.cell(row=r, column=1, value="By course (sections and students: all sections; ratings: see scope above)").font = f(True)
    hdr(r + 1, ["Course", "Sections", "Students taught", "Respondents (rated)", "Avg. Overall"])
    courses = sorted({x["course"] for x in recs})
    for i, cname in enumerate(courses):
        rr = r + 2 + i
        cond = f'({CRS}="{cname}")'
        label = f"{cname} — {titles[cname]}" if cname in titles else cname
        put(rr, [label, f"=SUMPRODUCT({cond}*1)", f"=SUMPRODUCT({cond}*{E_})", rsp(cond), wavg(cond)],
            [None, "0", "#,##0", "#,##0", "0.00"])
    r = r + 2 + len(courses)
    if summ_survey:
        S.cell(row=r, column=1, value=("Courses taught only under another survey version count in sections and students "
                                       "but show n/a for ratings.")).font = f(i=True, c=GRAY, s=9)
    r += 2
    S.cell(row=r, column=1, value="By year (sections and students: all sections; ratings: see scope above)").font = f(True)
    hdr(r + 1, ["Year", "Sections", "Students taught", "Respondents (rated)", "Avg. Overall"])
    years = sorted({x["year"] for x in recs})
    for i, y in enumerate(years):
        rr = r + 2 + i
        cond = f"({YR}={y})"
        put(rr, [y, f"=SUMPRODUCT({cond}*1)", f"=SUMPRODUCT({cond}*{E_})", rsp(cond), wavg(cond)],
            [None, "0", "#,##0", "#,##0", "0.00"])
    r = r + 2 + len(years) + 1

    if len(surveys) > 1:
        S.cell(row=r, column=1, value="By survey version (each version's own teaching items)").font = f(True)
        hdr(r + 1, ["Survey version", "Sections", "Respondents", "Avg. Overall"])
        for i, s in enumerate(surveys):
            rr = r + 2 + i
            cond = f'({SV}="{s.get("label", s["id"])}")'
            put(rr, [s.get("label", s["id"]), f"=SUMPRODUCT({cond}*{INC})", f"=SUMPRODUCT({cond}*{INC}*{RS})",
                     f'=IFERROR(SUMPRODUCT({cond}*{INC}*{OV}*{RS})/SUMPRODUCT({cond}*{INC}*{RS}),"n/a")'],
                [None, "0", "#,##0", "0.00"])
        r = r + 2 + len(surveys) + 1

    career_groups = scfg.get("career_avg_groups", primary)
    career_items = [it for it in items if it["group"] in career_groups]
    S.cell(row=r, column=1, value=f"3. Career average by evaluation item ({scope}, weighted by respondents)").font = f(True, c=NAVY, s=12)
    hdr(r + 1, ["Evaluation item", "", "", "Career avg"])
    S.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3)
    for j, it in enumerate(career_items):
        rr = r + 2 + j
        colr = rg(col[it["key"]])
        S.cell(row=rr, column=1, value=it["label"])
        S.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
        # only sections that actually have this item contribute
        S.cell(row=rr, column=4, value=(f'=IFERROR(SUMPRODUCT({SF}*{INC}*{RS}*({colr}<>"")*{colr})/'
                                        f'SUMPRODUCT({SF}*{INC}*{RS}*({colr}<>"")),"n/a")'))
        for cc in range(1, 5):
            S.cell(row=rr, column=cc).font, S.cell(row=rr, column=cc).border = f(), BORD
        S.cell(row=rr, column=4).number_format = "0.00"
        S.cell(row=rr, column=4).alignment = Alignment(horizontal="center")
    r = r + 2 + len(career_items) + 1

    top_groups = scfg.get("top_box_groups", [])
    top_items = [it for it in items if it["group"] in top_groups]
    if top_items:
        pooled = [x for x in recs if not summ_survey or x["survey"] == summ_survey]
        S.cell(row=r, column=1, value=f"4. STATIC — % of respondents choosing the top two options ({len(pooled)} sections pooled)").font = f(True, c=NAVY, s=12)
        S.cell(row=r + 1, column=1, value=("Computed once from the raw response counts on the reports (not from 'Course Index' columns). "
                                            "Denominator counts only the rating options, so 'Does Not Apply' is excluded.")).font = f(i=True, c=GRAY, s=9)
        hdr(r + 2, ["Evaluation item", "", "% top-two", "Ratings counted"])
        S.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=2)
        for j, it in enumerate(top_items):
            num = den = 0
            n_opt = len(it.get("scales", cfg.get("default_scales", [[5, 4, 3, 2, 1]]))[0])
            for x in pooled:
                v = x["items"].get(it["key"])
                if v and v.get("counts"):
                    cts = v["counts"][:n_opt]
                    num += sum(cts[:2])
                    den += sum(cts)
            rr = r + 3 + j
            S.cell(row=rr, column=1, value=it["label"])
            S.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=2)
            S.cell(row=rr, column=3, value=(num / den) if den else "n/a")
            S.cell(row=rr, column=4, value=den)
            for cc in range(1, 5):
                S.cell(row=rr, column=cc).font, S.cell(row=rr, column=cc).border = f(), BORD
            S.cell(row=rr, column=3).number_format = "0.0%"
            S.cell(row=rr, column=4).number_format = "#,##0"
            S.cell(row=rr, column=3).alignment = Alignment(horizontal="center")
            S.cell(row=rr, column=4).alignment = Alignment(horizontal="center")
    S.column_dimensions["A"].width = 52
    for cl, w in zip("BCDEF", (18, 16, 18, 20, 20)):
        S.column_dimensions[cl].width = w
    S.sheet_view.zoomScale = 95

    # ================= Verbatims =================
    V = wb.create_sheet("Student Verbatims")
    V.sheet_properties.tabColor = PURP
    V["A1"] = "Student Verbatims to Highlight"
    V["A1"].font = f(True, c=NAVY, s=14)
    V["A2"] = ("Exact student wording from evaluation reports and other feedback. Typos are preserved and student names removed. "
               "Filter by Theme or Priority ('Top pick').")
    V["A2"].font = f(i=True, c=GRAY, s=9)
    V.merge_cells("A2:J2")
    V["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    V.row_dimensions[2].height = 30
    vh = ["#", "Priority", "Theme", "Verbatim", "Course", "Term", "Format", "UG or GR", "Source type", "Source file"]
    vw = [5, 10, 26, 95, 14, 18, 12, 9, 30, 40]
    for i, (h, w) in enumerate(zip(vh, vw), 1):
        cell = V.cell(row=4, column=i, value=h)
        cell.font, cell.fill, cell.border = f(True, c="FFFFFF"), fill(NAVY), BORD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        V.column_dimensions[L(i)].width = w
    verbatims = verbatims or []
    theme_ord = cfg.get("verbatim_theme_order") or list(dict.fromkeys(q["theme"] for q in verbatims))
    lookup = {(x["course"], x["term_label"]): x for x in recs}
    rows = sorted(verbatims, key=lambda q: (theme_ord.index(q["theme"]) if q["theme"] in theme_ord else 99,
                                            0 if q.get("top_pick") else 1))
    for i, q in enumerate(rows, 1):
        r = 4 + i
        sec = lookup.get((q.get("course", ""), q.get("term", "")), {})
        vals = [i, "Top pick" if q.get("top_pick") else "", q["theme"], "“" + q["quote"] + "”", q.get("course", ""),
                q.get("term", ""), q.get("format") or sec.get("format", ""), q.get("level") or sec.get("level", ""),
                q.get("source_type", ""), q.get("source_file", "")]
        for cc, v in enumerate(vals, 1):
            cell = V.cell(row=r, column=cc, value=v)
            cell.font, cell.border = f(b=(cc == 2 and bool(v))), BORD
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if cc in (1, 2, 7, 8) else "left")
        if q.get("top_pick"):
            for cc in range(1, 11):
                V.cell(row=r, column=cc).fill = fill("E2EFDA")
        V.row_dimensions[r].height = max(30, 15 * (len(q["quote"]) // 85 + 1))
    V.freeze_panes = "D5"
    V.auto_filter.ref = f"A4:J{4 + max(len(rows), 1)}"

    # ================= Method & Notes =================
    M = wb.create_sheet("Method & Notes")
    M.sheet_properties.tabColor = GRAY
    M.column_dimensions["A"].width = 30
    M.column_dimensions["B"].width = 120
    M["A1"] = "Method & Notes"
    M["A1"].font = f(True, c=NAVY, s=14)
    online_rx = cfg.get("sections", {}).get("online_code_regex", "D")
    avg_items = [it["label"] for it in items if it["group"] in avg_groups]
    auto = [
        ("Source", f"{n_rows} course-evaluation reports, one per section or combined report. Only sections with a report on file are included."),
        ("Format", f"'Online' if any section code matches the pattern /{online_rx}/; otherwise 'In-person'. "
                   "For combined reports, the report is Online if any listed section code matches."),
        ("UG or GR", f"GR = course number {cfg.get('sections', {}).get('graduate_min_course_number', 500)} or higher; all others UG."),
        ("# taught", "Enrolled students on the report (not the number who responded)."),
        ("Category ratings", "The mean printed on each report for each survey item. Where a report prints only response counts, "
                             "the mean is recomputed from the counts (validated against printed means on the other reports)."),
        ("Avg. Overall", f"Simple (unweighted) average of the item means in the groups flagged in_avg_overall ({len(avg_items)} items). "
                         "Evaluations do not print a composite, so this is a calculated summary; every input is visible on the same row."),
        ("Summary averages", f"Respondent-weighted (section Avg. Overall × respondents). Sections with fewer than {small_n} respondents are flagged "
                             f"yellow; they are counted in summary averages when respondents >= {min_resp}."),
        ("Verbatims", "See 'Student Verbatims.' Wording is exact; names are removed."),
    ]
    extra = cfg.get("method_notes", [])
    for i, (a, b) in enumerate(auto + [tuple(e) for e in extra]):
        r = 3 + i
        M.cell(row=r, column=1, value=a).font = f(True)
        M.cell(row=r, column=2, value=b).font = f()
        for cc in (1, 2):
            M.cell(row=r, column=cc).alignment = Alignment(wrap_text=True, vertical="top")
            M.cell(row=r, column=cc).border = BORD
        M.row_dimensions[r].height = max(18, 15 * (len(b) // 115 + 1))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    if not quiet:
        print(f"saved {out_path}: {n_rows} sections, rows {R0}-{R1}")
    return out_path


def load_verbatims(path):
    if not path or not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    for r in rows:
        r["top_pick"] = str(r.get("top_pick", "")).strip().lower() in ("1", "y", "yes", "true", "x", "top pick")
    return [r for r in rows if r.get("quote")]
